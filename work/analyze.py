import re
from statistics import mean, median
from openpyxl import load_workbook
from pypdf import PdfReader

XLSX = "work/source/2026_actuals.xlsx"
PDF = "work/source/2025_final.pdf"

codes = {
    "서울": "C5S12", "성남": "C5S84", "청주": "C5V96", "구미": "C5V06",
    "인천북부": "C5S72", "광주": "C5V34", "부천": "C5S99", "서울동부": "C5S55",
    "서울남부": "C5S29", "구로": "C5T58", "대구동부": "C5U64", "부산동부": "C5U33",
    "고양": "C5T49", "전주": "C5U71", "천안": "C5W04", "서영사무소": "C5V37",
}

match_2025 = {
    "성남": "C5O27", "구미": "C5O37", "인천북부": "C5O43", "광주": "C5O30",
    "부천": "C5O26", "서울동부": "C5O34", "서울남부": "C5M24", "구로": "C5O25",
    "대구동부": "C5L76", "부산동부": "C5O32", "고양": "C5O29", "전주": "C5O35",
    "천안": "C5O36",
}

wb = load_workbook(XLSX, data_only=True)
summary = wb["요약"]
raw = wb["종합"]
headers = {raw.cell(1, c).value: c for c in range(1, 62)}
by_code = {raw.cell(r, headers["기관코드"]).value: r for r in range(2, 500) if raw.cell(r, headers["기관코드"]).value}

factors = [
    ("취업", "취업률", "취업점수환산"),
    ("알선", "알선취업률", "알선점수환산"),
    ("조기", "조기취업률", "조기점수환산"),
    ("고임금", "고임금률", "고임금점수환산"),
    ("유지", "고용유지율", "고용유지표준점수"),
]

all_rows = list(range(2, 289))
national_medians = {
    label: median(float(raw.cell(r, headers[score]).value or 0) for r in all_rows)
    for label, _, score in factors
}

def percentile(values, value):
    return 100 * sum(v <= value for v in values) / len(values)

data = []
for sr in range(7, 23):
    center = summary.cell(sr, 3).value
    rr = by_code[codes[center]]
    rec = {
        "center": center,
        "type": summary.cell(sr, 4).value,
        "june_score": float(summary.cell(sr, 5).value),
        "aug_score": float(summary.cell(sr, 6).value),
        "june_rank": int(summary.cell(sr, 7).value),
        "aug_rank": int(summary.cell(sr, 8).value),
        "june_grade": summary.cell(sr, 9).value,
        "aug_grade": summary.cell(sr, 10).value,
        "target": summary.cell(sr, 11).value,
        "risk": summary.cell(sr, 14).value,
        "focus": summary.cell(sr, 15).value,
        "action": summary.cell(sr, 17).value,
        "rates": {}, "scores": {}, "percentiles": {},
    }
    for label, rate, score in factors:
        rec["rates"][label] = 100 * float(raw.cell(rr, headers[rate]).value or 0)
        rec["scores"][label] = float(raw.cell(rr, headers[score]).value or 0)
        vals = [float(raw.cell(r, headers[score]).value or 0) for r in all_rows]
        rec["percentiles"][label] = percentile(vals, rec["scores"][label])
    data.append(rec)

text = "\n".join((p.extract_text() or "") for p in PdfReader(PDF).pages)
pdf_rows = {}
pat = re.compile(
    r"^(\d+)\s+(C5\w+)\s+(.+?)\s+(단독|컨소)\s+"
    r"([0-9.]+)\s+([0-9.]+)\s+([0-9.]+)\s+([0-9.]+)\s+([0-9.]+)\s+([0-9.]+)\s+"
    r"(?:(0\.5|1|2)\s+)?([0-9.]+)\s+(\d+)\s+([ABCD])$"
)
for line in text.splitlines():
    m = pat.match(line.strip())
    if not m:
        continue
    rank0, code, name, kind, *rest = m.groups()
    six = list(map(float, rest[:6]))
    bonus = float(rest[6]) if rest[6] else 0.0
    total, rank, grade = float(rest[7]), int(rest[8]), rest[9]
    pdf_rows[code] = {"name": name, "type": kind, "scores": six[:5], "satisfaction": six[5], "bonus": bonus, "total": total, "rank": rank, "grade": grade}

grade_floor = {"A": 45.9929350079649, "B": 43.6690651278641, "C": 42.3585665458295, "D": 38.1133134316125}

print("NATIONAL_MEDIANS", national_medians)
print("PORTFOLIO", {
    "avg_score_june": mean(x["june_score"] for x in data),
    "avg_score_aug": mean(x["aug_score"] for x in data),
    "avg_rank_june": mean(x["june_rank"] for x in data),
    "avg_rank_aug": mean(x["aug_rank"] for x in data),
    "grades_june": {g: sum(x["june_grade"] == g for x in data) for g in "ABCD"},
    "grades_aug": {g: sum(x["aug_grade"] == g for x in data) for g in "ABCD"},
    "score_up": sum(x["aug_score"] > x["june_score"] for x in data),
    "rank_up": sum(x["aug_rank"] < x["june_rank"] for x in data),
})

print("CENTER_ROWS")
for x in data:
    gap = max(0, grade_floor[x["target"]] - x["aug_score"])
    weakest = sorted(factors, key=lambda f: x["percentiles"][f[0]])[:2]
    y25 = pdf_rows.get(match_2025.get(x["center"], ""))
    deltas = None
    if y25:
        deltas = {label: x["scores"][label] - y25["scores"][i] for i, (label, _, _) in enumerate(factors)}
    print({
        "center": x["center"], "type": x["type"],
        "score": f'{x["june_score"]:.3f}->{x["aug_score"]:.3f} ({x["aug_score"]-x["june_score"]:+.3f})',
        "rank": f'{x["june_rank"]}->{x["aug_rank"]} ({x["june_rank"]-x["aug_rank"]:+d})',
        "grade": f'{x["june_grade"]}->{x["aug_grade"]}', "target": x["target"], "gap": round(gap,3),
        "rates": {k: round(v,1) for k,v in x["rates"].items()},
        "scores": {k: round(v,3) for k,v in x["scores"].items()},
        "weakest_pct": [(f[0], round(x["percentiles"][f[0]],1)) for f in weakest],
        "y25": None if not y25 else f'{y25["grade"]}/{y25["rank"]}/{y25["total"]:.3f}',
        "delta25": None if deltas is None else {k: round(v,3) for k,v in deltas.items()},
        "risk": x["risk"], "focus": x["focus"], "action": x["action"],
    })

print("MATCHED_2025", len(match_2025), "PARSED", len(pdf_rows))

print("PORTFOLIO_FACTORS", {
    "rate_avg": {label: round(mean(x["rates"][label] for x in data), 1) for label, _, _ in factors},
    "score_avg": {label: round(mean(x["scores"][label] for x in data), 3) for label, _, _ in factors},
    "employment_60_count": sum(x["rates"]["취업"] >= 60 for x in data),
    "placement_10_count": sum(x["rates"]["알선"] >= 10 for x in data),
    "both_count": sum(x["rates"]["취업"] >= 60 and x["rates"]["알선"] >= 10 for x in data),
})
matched_deltas = []
for x in data:
    y25 = pdf_rows.get(match_2025.get(x["center"], ""))
    if y25:
        matched_deltas.append({label: x["scores"][label] - y25["scores"][i] for i, (label, _, _) in enumerate(factors)})
print("AVG_DELTA_2025_TO_2026_AUG", {label: round(mean(d[label] for d in matched_deltas), 3) for label, _, _ in factors})
print("TARGET_GAP", {x["center"]: round(max(0, grade_floor[x["target"]] - x["aug_score"]), 3) for x in data})
